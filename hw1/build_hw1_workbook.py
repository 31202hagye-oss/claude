"""Build hw1_frontier.xlsx: mean-SD and mean-variance frontiers for HW1, Problem 2.

All numbers in the workbook are live Excel formulas driven by the Parameters sheet,
so changing any input (e.g. rf) updates every table and chart.

    python3 hw1/build_hw1_workbook.py

If LibreOffice is installed, the file is recalculated and re-saved through it so the
workbook also stores computed values; viewers that do not recalculate (phone/app
previews) can then display the numbers and charts.
"""
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.marker import Marker
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).with_name("hw1_frontier.xlsx")

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")

# Parameter cell addresses (sheet "Parameters")
P = {
    "mu1": "Parameters!$B$4", "mu2": "Parameters!$B$5",
    "s1": "Parameters!$B$6", "s2": "Parameters!$B$7",
    "rho": "Parameters!$B$8", "rf": "Parameters!$B$9",
    "s12": "Parameters!$B$12", "e1": "Parameters!$B$13", "e2": "Parameters!$B$14",
    "w": "Parameters!$B$15", "w2": "Parameters!$B$16",
    "mup": "Parameters!$B$17", "varp": "Parameters!$B$18", "sdp": "Parameters!$B$19",
    "sr": "Parameters!$B$20",
    "wg": "Parameters!$B$22", "mug": "Parameters!$B$23", "varg": "Parameters!$B$24",
    "a": "Parameters!$B$26",
    "ray": "Parameters!$B$28", "tan": "Parameters!$B$29",
}


def header(ws, row, labels):
    for col, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=col, value=label)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def build_parameters(wb):
    ws = wb.active
    ws.title = "Parameters"
    ws["A1"] = "HW1 – Two risky assets + one risk-free asset"
    ws["A1"].font = TITLE
    ws["A3"] = "Inputs (yellow cells can be changed)"
    ws["A3"].font = BOLD
    inputs = [
        ("μ1", 0.10), ("μ2", 0.14), ("σ1", 0.15), ("σ2", 0.20), ("ρ12", 0.2),
        ("rf", 0.05),
    ]
    for i, (name, val) in enumerate(inputs, start=4):
        ws.cell(row=i, column=1, value=name)
        c = ws.cell(row=i, column=2, value=val)
        c.fill = INPUT_FILL
    ws["C9"] = "rf is not listed in the problem; 5% assumed – change here if your class used another value."

    ws["A11"] = "Derived quantities"
    ws["A11"].font = BOLD
    derived = [
        (12, "σ12 = ρ12·σ1·σ2", "=B8*B6*B7"),
        (13, "μe1 = μ1 − rf", "=B4-B9"),
        (14, "μe2 = μ2 − rf", "=B5-B9"),
        (15, "w (weight on asset 1, Eq. 1)",
         "=(B13*B7^2-B14*B12)/(B13*B7^2+B14*B6^2-(B13+B14)*B12)"),
        (16, "1 − w (weight on asset 2)", "=1-B15"),
        (17, "μp", "=B15*B4+B16*B5"),
        (18, "σp²", "=B15^2*B6^2+B16^2*B7^2+2*B15*B16*B12"),
        (19, "σp", "=SQRT(B18)"),
        (20, "Sharpe ratio (μp − rf)/σp", "=(B17-B9)/B19"),
        (22, "w_GMV (global min-variance)", "=(B7^2-B12)/(B6^2+B7^2-2*B12)"),
        (23, "μ_GMV", "=B22*B4+(1-B22)*B5"),
        (24, "σ²_GMV", "=B22^2*B6^2+(1-B22)^2*B7^2+2*B22*(1-B22)*B12"),
        (25, "σ_GMV", "=SQRT(B24)"),
        (26, "a = (σ1²+σ2²−2σ12)/(μ2−μ1)²", "=(B6^2+B7^2-2*B12)/(B5-B4)^2"),
        (27, "Analytic parabola:", None),
        (28, "Ray slope (μp − rf)/σp²", "=(B17-B9)/B18"),
        (29, "Tangent slope at p = (μp − rf)/(2σp²)", "=(B17-B9)/(2*B18)"),
        (30, "Tangent slope ÷ ray slope", "=B29/B28"),
    ]
    for row, label, formula in derived:
        ws.cell(row=row, column=1, value=label)
        if formula:
            ws.cell(row=row, column=2, value=formula)
    ws["B27"] = "σ² = σ²_GMV + a·(μ − μ_GMV)²"
    for row in list(range(4, 10)) + list(range(12, 27)) + [28, 29, 30]:
        ws.cell(row=row, column=2).number_format = "0.000000"
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 16


def build_frontier(wb):
    ws = wb.create_sheet("Frontier")
    ws["A1"] = "Risky-asset frontier for w1 = −0.5 … 1.5 (step 0.05)"
    ws["A1"].font = BOLD
    header(ws, 2, ["w1", "w2 = 1 − w1", "μc", "σc", "σc² (from weights)",
                   "σc² (analytic parabola)", "difference"])
    w_rows = 41
    for i in range(w_rows):
        r = 3 + i
        ws.cell(row=r, column=1, value=round(-0.5 + 0.05 * i, 2))
        ws.cell(row=r, column=2, value=f"=1-A{r}")
        ws.cell(row=r, column=3, value=f"=A{r}*{P['mu1']}+B{r}*{P['mu2']}")
        ws.cell(row=r, column=5, value=(
            f"=A{r}^2*{P['s1']}^2+B{r}^2*{P['s2']}^2+2*A{r}*B{r}*{P['s12']}"))
        ws.cell(row=r, column=4, value=f"=SQRT(E{r})")
        ws.cell(row=r, column=6, value=f"={P['varg']}+{P['a']}*(C{r}-{P['mug']})^2")
        ws.cell(row=r, column=7, value=f"=E{r}-F{r}")
        ws.cell(row=r, column=7).number_format = "0.0E+00"
    last_w = 2 + w_rows

    # Frontier with the risk-free asset, mean-SD space (two straight lines)
    ws["I1"] = "With rf: mean–SD (CML)"
    ws["I1"].font = BOLD
    header_cols = {9: "σc", 10: "μc upper = rf + SR·σc", 11: "μc lower = rf − SR·σc"}
    for col, label in header_cols.items():
        c = ws.cell(row=2, column=col, value=label)
        c.font, c.fill = BOLD, HEAD_FILL
    sd_rows = 31
    for i in range(sd_rows):
        r = 3 + i
        ws.cell(row=r, column=9, value=round(0.01 * i, 2))
        ws.cell(row=r, column=10, value=f"={P['rf']}+{P['sr']}*I{r}")
        ws.cell(row=r, column=11, value=f"={P['rf']}-{P['sr']}*I{r}")
    last_sd = 2 + sd_rows

    # Frontier with the risk-free asset, mean-variance space (Eq. 2)
    ws["M1"] = "With rf: mean–variance (Eq. 2)"
    ws["M1"].font = BOLD
    header_cols = {13: "σc²", 14: "μc upper = rf + SR·√σc²", 15: "μc lower = rf − SR·√σc²",
                   16: "Ray rf→p: rf + [(μp−rf)/σp²]·σc²",
                   17: "Tangent at p: μp + [(μp−rf)/(2σp²)]·(σc²−σp²)"}
    for col, label in header_cols.items():
        c = ws.cell(row=2, column=col, value=label)
        c.font, c.fill = BOLD, HEAD_FILL
    var_rows = 46
    for i in range(var_rows):
        r = 3 + i
        ws.cell(row=r, column=13, value=round(0.002 * i, 3))
        ws.cell(row=r, column=14, value=f"={P['rf']}+{P['sr']}*SQRT(M{r})")
        ws.cell(row=r, column=15, value=f"={P['rf']}-{P['sr']}*SQRT(M{r})")
        ws.cell(row=r, column=16, value=f"={P['rf']}+{P['ray']}*M{r}")
        ws.cell(row=r, column=17, value=f"={P['mup']}+{P['tan']}*(M{r}-{P['varp']})")
    last_var = 2 + var_rows

    # Analytic parabola on a μ grid
    ws["S1"] = "Analytic parabola on a μ grid"
    ws["S1"].font = BOLD
    for col, label in {19: "μc", 20: "σc² = σ²_GMV + a(μc − μ_GMV)²", 21: "σc"}.items():
        c = ws.cell(row=2, column=col, value=label)
        c.font, c.fill = BOLD, HEAD_FILL
    mu_rows = 33
    for i in range(mu_rows):
        r = 3 + i
        ws.cell(row=r, column=19, value=round(0.08 + 0.0025 * i, 4))
        ws.cell(row=r, column=20, value=f"={P['varg']}+{P['a']}*(S{r}-{P['mug']})^2")
        ws.cell(row=r, column=21, value=f"=SQRT(T{r})")
    last_mu = 2 + mu_rows

    # Key points
    ws["W1"] = "Key points"
    ws["W1"].font = BOLD
    header_cells = {23: "Point", 24: "σ", 25: "σ²", 26: "μ"}
    for col, label in header_cells.items():
        c = ws.cell(row=2, column=col, value=label)
        c.font, c.fill = BOLD, HEAD_FILL
    points = [
        ("Asset 1", f"={P['s1']}", f"={P['mu1']}"),
        ("Asset 2", f"={P['s2']}", f"={P['mu2']}"),
        ("Tangency p", f"={P['sdp']}", f"={P['mup']}"),
        ("Risk-free", "=0", f"={P['rf']}"),
        ("GMV", f"=SQRT({P['varg']})", f"={P['mug']}"),
    ]
    for i, (name, sd, mu) in enumerate(points):
        r = 3 + i
        ws.cell(row=r, column=23, value=name)
        ws.cell(row=r, column=24, value=sd)
        ws.cell(row=r, column=25, value=f"=X{r}^2")
        ws.cell(row=r, column=26, value=mu)
    ws["W9"] = "Tangency weights:"
    ws["W9"].font = BOLD
    ws["W10"] = "w1 (asset 1)"
    ws["X10"] = f"={P['w']}"
    ws["W11"] = "w2 (asset 2)"
    ws["X11"] = f"={P['w2']}"

    for col in "ABCDEFIJKMNOPQSTUXYZ":
        ws.column_dimensions[col].width = 13
    ws.column_dimensions["W"].width = 16
    return ws, last_w, last_sd, last_var, last_mu


def line_series(ws, xcol, ycol, first, last, title, color, dash=None, width=22000):
    s = Series(Reference(ws, min_col=ycol, min_row=first, max_row=last),
               Reference(ws, min_col=xcol, min_row=first, max_row=last), title=title)
    s.marker = Marker(symbol="none")
    s.smooth = True
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = width
    if dash:
        s.graphicalProperties.line.dashStyle = dash
    return s


def point_series(ws, xcol, ycol, row, title, color, symbol="circle", size=9):
    s = Series(Reference(ws, min_col=ycol, min_row=row, max_row=row),
               Reference(ws, min_col=xcol, min_row=row, max_row=row), title=title)
    s.marker = Marker(symbol=symbol, size=size)
    s.marker.graphicalProperties.solidFill = color
    s.marker.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.noFill = True
    return s


def make_chart(title, xtitle, ytitle, xmax, ymin=0.0, ymax=0.22):
    ch = ScatterChart()
    ch.title = title
    ch.style = 13
    ch.x_axis.title = xtitle
    ch.y_axis.title = ytitle
    ch.x_axis.scaling.min = 0
    ch.x_axis.scaling.max = xmax
    ch.y_axis.scaling.min = ymin
    ch.y_axis.scaling.max = ymax
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.majorGridlines = None
    ch.height = 13
    ch.width = 24
    ch.legend.position = "b"
    return ch


def add_points(ch, ws, xcol):
    ch.series.append(point_series(ws, xcol, 26, 3, "Asset 1", "70AD47", "square"))
    ch.series.append(point_series(ws, xcol, 26, 4, "Asset 2", "7030A0", "square"))
    ch.series.append(point_series(ws, xcol, 26, 5, "Tangency p", "C00000", "diamond", 11))
    ch.series.append(point_series(ws, xcol, 26, 6, "Risk-free", "000000", "circle"))
    ch.series.append(point_series(ws, xcol, 26, 7, "GMV", "808080", "triangle"))


def build_charts(wb, ws, last_w, last_sd, last_var, last_mu):
    cs = wb.create_sheet("Charts")
    cs["A1"] = "Charts (data on the Frontier sheet)"
    cs["A1"].font = TITLE
    cs["A2"] = "Tangency weights:  w1 ="
    cs["C2"] = f"={P['w']}"
    cs["D2"] = "w2 ="
    cs["E2"] = f"={P['w2']}"
    cs["F2"] = "μp ="
    cs["G2"] = f"={P['mup']}"
    cs["H2"] = "σp ="
    cs["I2"] = f"={P['sdp']}"
    for c in ("C2", "E2", "G2", "I2"):
        cs[c].number_format = "0.0000"

    # 1. Mean-SD frontier with and without rf
    ch1 = make_chart("Mean–standard deviation frontier", "Standard deviation σ",
                     "Expected return μ", 0.30, ymin=-0.05)
    ch1.series.append(line_series(ws, 4, 3, 3, last_w, "Risky assets only", "4472C4"))
    ch1.series.append(line_series(ws, 9, 10, 3, last_sd, "With rf (upper CML)", "ED7D31"))
    ch1.series.append(line_series(ws, 9, 11, 3, last_sd, "With rf (lower)", "ED7D31",
                                  dash="dash"))
    add_points(ch1, ws, 24)
    cs.add_chart(ch1, "A4")

    # 2. Mean-variance frontier with and without rf (Eq. 2), simulated vs analytic
    ch2 = make_chart("Mean–variance frontier", "Variance σ²", "Expected return μ", 0.09,
                     ymin=-0.05)
    sim = Series(Reference(ws, min_col=3, min_row=3, max_row=last_w),
                 Reference(ws, min_col=5, min_row=3, max_row=last_w),
                 title="Risky assets: from weights (w1 grid)")
    sim.marker = Marker(symbol="circle", size=5)
    sim.marker.graphicalProperties.solidFill = "4472C4"
    sim.marker.graphicalProperties.line.solidFill = "4472C4"
    sim.graphicalProperties.line.noFill = True
    ch2.series.append(sim)
    ch2.series.append(line_series(ws, 20, 19, 3, last_mu, "Risky assets: analytic parabola",
                                  "1F3864", width=12700))
    ch2.series.append(line_series(ws, 13, 14, 3, last_var, "With rf: rf + SR·√σ² (Eq. 2)",
                                  "ED7D31"))
    ch2.series.append(line_series(ws, 13, 15, 3, last_var, "With rf: rf − SR·√σ²",
                                  "ED7D31", dash="dash"))
    add_points(ch2, ws, 25)
    cs.add_chart(ch2, "A31")

    # 3. Ray vs tangent in mean-variance space
    ch3 = make_chart("Ray rf→p vs tangent at p (mean–variance space)", "Variance σ²",
                     "Expected return μ", 0.06, ymin=0.0, ymax=0.30)
    ch3.series.append(line_series(ws, 20, 19, 3, last_mu, "Risky-asset frontier", "4472C4"))
    ch3.series.append(line_series(ws, 13, 14, 3, last_var, "Eq. 2 frontier", "ED7D31"))
    ch3.series.append(line_series(ws, 13, 16, 3, last_var, "Ray rf→p (slope (μp−rf)/σp²)",
                                  "C00000", dash="dash", width=15000))
    ch3.series.append(line_series(ws, 13, 17, 3, last_var,
                                  "Tangent at p (slope (μp−rf)/(2σp²))", "00B050",
                                  dash="sysDot", width=15000))
    add_points(ch3, ws, 25)
    cs.add_chart(ch3, "A58")


def build_slopes(wb):
    ws = wb.create_sheet("Slopes")
    ws["A1"] = "Ray slope vs tangent slope at points on the risky-asset frontier (σ² on x-axis)"
    ws["A1"].font = BOLD
    ws["A2"] = ("Tangent slope dμ/dσ² = (μ1 − μ2) / (2[w·σ1² − (1−w)·σ2² + (1−2w)·σ12]);  "
                "ray slope = (μ − rf)/σ².  Only at p is tangent = ½ × ray.")
    header(ws, 4, ["w1", "μ", "σ²", "Ray slope (μ − rf)/σ²", "Tangent slope dμ/dσ²",
                   "Tangent ÷ ray", "Tangent-line intercept at σ²=0",
                   "(rf + μ)/2", "Point"])
    ws_list = [-0.5, -0.2, 0.0, 0.1, 0.2, 0.3, 0.4, "p", 0.5, 0.6, 0.65, 1.0]
    for i, w in enumerate(ws_list):
        r = 5 + i
        if w == "p":
            ws.cell(row=r, column=1, value=f"={P['w']}")
            ws.cell(row=r, column=9, value="Tangency p")
            for col in range(1, 10):
                ws.cell(row=r, column=col).font = BOLD
        else:
            ws.cell(row=r, column=1, value=w)
            note = "inefficient (below GMV)" if w >= 0.6733 else ""
            ws.cell(row=r, column=9, value=note)
        ws.cell(row=r, column=2, value=f"=A{r}*{P['mu1']}+(1-A{r})*{P['mu2']}")
        ws.cell(row=r, column=3, value=(
            f"=A{r}^2*{P['s1']}^2+(1-A{r})^2*{P['s2']}^2+2*A{r}*(1-A{r})*{P['s12']}"))
        ws.cell(row=r, column=4, value=f"=(B{r}-{P['rf']})/C{r}")
        ws.cell(row=r, column=5, value=(
            f"=({P['mu1']}-{P['mu2']})/(2*(A{r}*{P['s1']}^2-(1-A{r})*{P['s2']}^2"
            f"+(1-2*A{r})*{P['s12']}))"))
        ws.cell(row=r, column=6, value=f"=E{r}/D{r}")
        ws.cell(row=r, column=7, value=f"=B{r}-E{r}*C{r}")
        ws.cell(row=r, column=8, value=f"=({P['rf']}+B{r})/2")
        for col in range(1, 9):
            ws.cell(row=r, column=col).number_format = "0.0000"
    for col, width in zip("ABCDEFGHI", (10, 10, 10, 16, 16, 12, 18, 12, 24)):
        ws.column_dimensions[col].width = width


def main():
    wb = Workbook()
    build_parameters(wb)
    ws, last_w, last_sd, last_var, last_mu = build_frontier(wb)
    build_charts(wb, ws, last_w, last_sd, last_var, last_mu)
    build_slopes(wb)
    wb.move_sheet("Charts", offset=-2)
    wb.save(OUT)
    embed_cached_values(OUT)
    print(f"wrote {OUT}")


def embed_cached_values(path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice not found: saved formulas only (Excel recalculates on open)")
        return
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                        "--outdir", tmp, str(path)], check=True, capture_output=True)
        shutil.move(Path(tmp) / path.name, path)


if __name__ == "__main__":
    main()
